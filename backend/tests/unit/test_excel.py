"""Unit tests for ExcelExporter with delta tracking and visual highlighting."""

import os
import tempfile

import openpyxl

from backend.app.exporters.excel import ExcelExporter
from backend.app.models.job import ExtractionStats, ProcessedJobRecord


def test_excel_export_structure() -> None:
    records = [
        ProcessedJobRecord(
            "ABC Technologies", "Software Engineer", "10 - 50 employees", is_new=True
        ),
        ProcessedJobRecord(
            "XYZ Solutions", "Backend Developer", "More than 1000 employees", is_new=False
        ),
        ProcessedJobRecord("Global Enterprises", "Product Lead", "N/A", is_new=True),
    ]
    stats = ExtractionStats(
        pages_processed=1,
        jobs_discovered=4,
        jobs_processed=3,
        companies_discovered=3,
        duplicates_removed=1,
        missing_employee_counts=1,
        new_jobs_added=2,
        existing_jobs_seen=1,
    )

    with tempfile.TemporaryDirectory() as temp_dir:
        output_file = os.path.join(temp_dir, "test_jobs.xlsx")
        result_path = ExcelExporter.export(
            records=records,
            stats=stats,
            source_url="https://example.com/careers",
            output_path=output_file,
        )

        assert os.path.exists(result_path)

        wb = openpyxl.load_workbook(result_path)

        # Check sheets
        assert "All Jobs" in wb.sheetnames
        assert "New Jobs Only" in wb.sheetnames
        assert "Existing Jobs Only" in wb.sheetnames
        assert "Small (1-50)" in wb.sheetnames
        assert "Medium (50-500)" in wb.sheetnames
        assert "Large (500+)" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        # Check All Jobs sheet headers
        ws_all = wb["All Jobs"]
        headers = [ws_all.cell(row=1, column=col).value for col in range(1, 6)]
        assert headers == ["Company Name", "Job Role", "Number of People", "Date Posted / Discovered", "Status"]

        # Check records on All Jobs
        row2 = [ws_all.cell(row=2, column=col).value for col in range(1, 6)]
        assert row2[0] == "ABC Technologies"
        assert row2[1] == "Software Engineer"
        assert row2[2] == "10 - 50 employees"
        assert row2[4] == "NEW"

        row3 = [ws_all.cell(row=3, column=col).value for col in range(1, 6)]
        assert row3[0] == "XYZ Solutions"
        assert row3[1] == "Backend Developer"
        assert row3[2] == "More than 1000 employees"
        assert row3[4] == "EXISTING"

        row4 = [ws_all.cell(row=4, column=col).value for col in range(1, 6)]
        assert row4[0] == "Global Enterprises"
        assert row4[1] == "Product Lead"
        assert row4[2] == "N/A"
        assert row4[4] == "NEW"

        # Check New Jobs Only sheet
        ws_new = wb["New Jobs Only"]
        assert ws_new.max_row == 3  # Header + 2 new records
        new_row1 = [ws_new.cell(row=2, column=col).value for col in range(1, 6)]
        assert new_row1[0] == "ABC Technologies"
        assert new_row1[4] == "NEW"

        # Check Existing Jobs Only sheet
        ws_existing = wb["Existing Jobs Only"]
        assert ws_existing.max_row == 2  # Header + 1 existing record
        exist_row1 = [ws_existing.cell(row=2, column=col).value for col in range(1, 6)]
        assert exist_row1[0] == "XYZ Solutions"
        assert exist_row1[4] == "EXISTING"


def test_export_category_workbook() -> None:
    category_meta = {
        "id": "qa_automation",
        "name": "QA & Automation",
        "filename": "QA_Automation_Jobs.xlsx",
        "keywords": ["Selenium", "Playwright"],
    }

    rec_selenium_small = ProcessedJobRecord(
        company_name="Alpha QA",
        job_role="Senior Selenium Engineer",
        number_of_people="10 - 50 employees",
        is_new=True,
    )
    rec_selenium_big = ProcessedJobRecord(
        company_name="MegaCorp",
        job_role="Lead Selenium SDET",
        number_of_people="500 - 1000 employees",
        is_new=False,
    )
    rec_playwright_small = ProcessedJobRecord(
        company_name="Beta Labs",
        job_role="Playwright Automation Engineer",
        number_of_people="1 - 10 employees",
        is_new=True,
    )

    category_data = {
        "all_category_jobs": [rec_selenium_small, rec_selenium_big, rec_playwright_small],
        "all_under_100": [rec_selenium_small, rec_playwright_small],
        "keywords": {
            "Selenium": {
                "all": [rec_selenium_small, rec_selenium_big],
                "under_100": [rec_selenium_small],
            },
            "Playwright": {
                "all": [rec_playwright_small],
                "under_100": [rec_playwright_small],
            },
        },
    }

    with tempfile.TemporaryDirectory() as temp_dir:
        output_file = os.path.join(temp_dir, "QA_Automation_Jobs.xlsx")
        result_path = ExcelExporter.export_category_workbook(
            category_meta=category_meta,
            category_data=category_data,
            output_path=output_file,
        )

        assert os.path.exists(result_path)
        wb = openpyxl.load_workbook(result_path)

        # Expected sheet names:
        # 1. All QA & Automation Jobs
        # 2. < 100 People (All)
        # 3. Selenium
        # 4. Selenium (<100)
        # 5. Playwright
        # 6. Playwright (<100)
        # 7. Summary
        assert "All QA & Automation Jobs" in wb.sheetnames
        assert "< 100 People (All)" in wb.sheetnames
        assert "Selenium" in wb.sheetnames
        assert "Selenium (<100)" in wb.sheetnames
        assert "Playwright" in wb.sheetnames
        assert "Playwright (<100)" in wb.sheetnames
        assert "Summary" in wb.sheetnames

        # Verify < 100 People (All) has only small companies
        ws_u100 = wb["< 100 People (All)"]
        assert ws_u100.max_row == 3  # Header + 2 jobs
        u100_companies = [ws_u100.cell(row=r, column=1).value for r in range(2, 4)]
        assert "Alpha QA" in u100_companies
        assert "Beta Labs" in u100_companies
        assert "MegaCorp" not in u100_companies

        # Verify Selenium (<100) sheet has 1 job
        ws_sel_u100 = wb["Selenium (<100)"]
        assert ws_sel_u100.max_row == 2  # Header + 1 job
        assert ws_sel_u100.cell(row=2, column=1).value == "Alpha QA"

