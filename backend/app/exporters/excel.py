"""Excel generation component utilizing OpenPyXL with professional styling, delta highlighting, and pre-filtered category sheets."""

import os
import re
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from backend.app.core.exceptions import ExcelExportError
from backend.app.core.logging import get_logger
from backend.app.models.job import ExtractionStats, ProcessedJobRecord

logger = get_logger(__name__)


class ExcelExporter:
    """Generates structured, formatted Excel (.xlsx) workbooks with delta tracking and categorized tabs."""

    # Colors & Typography
    HEADER_FILL = PatternFill(
        start_color="1E293B", end_color="1E293B", fill_type="solid"
    )  # Slate 800
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    DATA_FONT = Font(name="Calibri", size=11, color="0F172A")
    SUMMARY_TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1E293B")
    SUMMARY_KEY_FONT = Font(name="Calibri", size=11, bold=True, color="475569")
    THIN_BORDER_SIDE = Side(border_style="thin", color="E2E8F0")
    THIN_BORDER = Border(
        left=THIN_BORDER_SIDE,
        right=THIN_BORDER_SIDE,
        top=THIN_BORDER_SIDE,
        bottom=THIN_BORDER_SIDE,
    )

    # Visual Delta Highlighting for New Records
    NEW_ROW_FILL = PatternFill(
        start_color="F0FDF4", end_color="F0FDF4", fill_type="solid"
    )  # Soft Emerald tint
    NEW_STATUS_FILL = PatternFill(
        start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"
    )  # Emerald badge
    NEW_STATUS_FONT = Font(name="Calibri", size=11, bold=True, color="166534")  # Dark Emerald
    EXISTING_STATUS_FONT = Font(name="Calibri", size=11, color="64748B")  # Slate Gray

    @classmethod
    def _populate_jobs_sheet(
        cls,
        ws,
        records: list[ProcessedJobRecord],
        table_name: str,
        empty_message: str = "No matching jobs in this category.",
    ) -> None:
        """Helper to render and format a table of jobs with Excel Table filters."""
        ws.views.sheetView[0].showGridLines = True
        headers = ["Company Name", "Job Role", "Number of People", "Date Posted / Discovered", "Status"]
        ws.append(headers)

        # Style header row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = cls.HEADER_FONT
            cell.fill = cls.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = cls.THIN_BORDER

        ws.row_dimensions[1].height = 26

        if not records:
            ws.append([empty_message, "", "", "", ""])
            ws.merge_cells("A2:E2")
            empty_cell = ws.cell(row=2, column=1)
            empty_cell.font = Font(name="Calibri", size=11, italic=True, color="64748B")
            empty_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 24
            return

        for row_idx, record in enumerate(records, start=2):
            row_data = record.to_row()
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20

            row_fill = cls.NEW_ROW_FILL if record.is_new else None

            # Company Name
            c1 = ws.cell(row=row_idx, column=1)
            c1.font = cls.DATA_FONT
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c1.border = cls.THIN_BORDER
            if row_fill:
                c1.fill = row_fill

            # Job Role
            c2 = ws.cell(row=row_idx, column=2)
            c2.font = cls.DATA_FONT
            c2.alignment = Alignment(horizontal="left", vertical="center")
            c2.border = cls.THIN_BORDER
            if row_fill:
                c2.fill = row_fill

            # Number of People
            c3 = ws.cell(row=row_idx, column=3)
            c3.font = cls.DATA_FONT
            c3.border = cls.THIN_BORDER
            if row_fill:
                c3.fill = row_fill
            if isinstance(record.number_of_people, (int, float)):
                c3.alignment = Alignment(horizontal="right", vertical="center")
                c3.number_format = "#,##0"
            elif str(record.number_of_people) == "N/A":
                c3.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c3.alignment = Alignment(horizontal="left", vertical="center")

            # Date Posted / Discovered
            c4 = ws.cell(row=row_idx, column=4)
            c4.font = cls.DATA_FONT
            c4.alignment = Alignment(horizontal="center", vertical="center")
            c4.border = cls.THIN_BORDER
            if row_fill:
                c4.fill = row_fill

            # Status (NEW vs EXISTING)
            c5 = ws.cell(row=row_idx, column=5)
            c5.border = cls.THIN_BORDER
            c5.alignment = Alignment(horizontal="center", vertical="center")
            if record.is_new:
                c5.font = cls.NEW_STATUS_FONT
                c5.fill = cls.NEW_STATUS_FILL
            else:
                c5.font = cls.EXISTING_STATUS_FONT

        # Setup Table with native Excel filters
        end_row = max(len(records) + 1, 2)
        safe_table_name = re.sub(r"[^A-Za-z0-9_]", "_", table_name)
        tab = Table(displayName=safe_table_name, ref=f"A1:E{end_row}")
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        try:
            ws.add_table(tab)
        except Exception:
            ws.auto_filter.ref = f"A1:E{end_row}"

        ws.freeze_panes = "A2"

        # Auto-size columns
        min_widths = {"A": 28, "B": 32, "C": 22, "D": 24, "E": 14}
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            calculated = max(max_len + 5, min_widths.get(col_letter, 18))
            ws.column_dimensions[col_letter].width = min(calculated, 65)


    @classmethod
    def export(
        cls,
        records: list[ProcessedJobRecord],
        stats: ExtractionStats,
        source_url: str,
        output_path: str,
    ) -> str:
        """
        Build and save the Excel workbook to the specified output path.
        Includes pre-filtered tabs for 1-click filtering:
        - Sheet 1: "All Jobs"
        - Sheet 2: "New Jobs Only"
        - Sheet 3: "Existing Jobs Only"
        - Sheet 4: "Small (1-50)"
        - Sheet 5: "Medium (50-500)"
        - Sheet 6: "Large (500+)"
        - Sheet 7: "Summary"
        """
        try:
            wb = openpyxl.Workbook()

            # ---------------------------------------------------------
            # Sheet 1: All Jobs
            # ---------------------------------------------------------
            ws_all = wb.active
            ws_all.title = "All Jobs"
            cls._populate_jobs_sheet(
                ws_all,
                records,
                table_name="Table_AllJobs",
                empty_message="No jobs extracted in this run.",
            )

            # ---------------------------------------------------------
            # Sheet 2: New Jobs Only
            # ---------------------------------------------------------
            new_records = [r for r in records if r.is_new]
            ws_new = wb.create_sheet(title="New Jobs Only")
            cls._populate_jobs_sheet(
                ws_new,
                new_records,
                table_name="Table_NewJobs",
                empty_message="No new jobs discovered in this run. All jobs were previously indexed.",
            )

            # ---------------------------------------------------------
            # Sheet 3: Existing Jobs Only
            # ---------------------------------------------------------
            existing_records = [r for r in records if not r.is_new]
            ws_existing = wb.create_sheet(title="Existing Jobs Only")
            cls._populate_jobs_sheet(
                ws_existing,
                existing_records,
                table_name="Table_ExistingJobs",
                empty_message="No existing jobs. All jobs in this run are newly discovered.",
            )

            # ---------------------------------------------------------
            # Sheet 4: Small Companies (1 - 50 employees)
            # ---------------------------------------------------------
            small_records = [
                r
                for r in records
                if "1 - 10" in str(r.number_of_people) or "10 - 50" in str(r.number_of_people)
            ]
            ws_small = wb.create_sheet(title="Small (1-50)")
            cls._populate_jobs_sheet(
                ws_small,
                small_records,
                table_name="Table_SmallSize",
                empty_message="No small companies (1 - 50 employees) in this dataset.",
            )

            # ---------------------------------------------------------
            # Sheet 5: Medium Companies (50 - 500 employees)
            # ---------------------------------------------------------
            medium_records = [
                r
                for r in records
                if "50 - 200" in str(r.number_of_people) or "200 - 500" in str(r.number_of_people)
            ]
            ws_medium = wb.create_sheet(title="Medium (50-500)")
            cls._populate_jobs_sheet(
                ws_medium,
                medium_records,
                table_name="Table_MediumSize",
                empty_message="No medium companies (50 - 500 employees) in this dataset.",
            )

            # ---------------------------------------------------------
            # Sheet 6: Large Companies (500+ employees)
            # ---------------------------------------------------------
            large_records = [
                r
                for r in records
                if "500 - 1000" in str(r.number_of_people)
                or "1000" in str(r.number_of_people)
                or "more than" in str(r.number_of_people).lower()
            ]
            ws_large = wb.create_sheet(title="Large (500+)")
            cls._populate_jobs_sheet(
                ws_large,
                large_records,
                table_name="Table_LargeSize",
                empty_message="No large enterprise companies (500+ employees) in this dataset.",
            )

            # ---------------------------------------------------------
            # Sheet 7: Summary
            # ---------------------------------------------------------
            ws_summary = wb.create_sheet(title="Summary")
            ws_summary.views.sheetView[0].showGridLines = True

            ws_summary.append(["EXTRACTION SUMMARY", ""])
            ws_summary.merge_cells("A1:B1")
            title_cell = ws_summary.cell(row=1, column=1)
            title_cell.font = cls.SUMMARY_TITLE_FONT
            title_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws_summary.row_dimensions[1].height = 30

            new_count = stats.new_jobs_added if stats.new_jobs_added else len(new_records)
            existing_count = (
                stats.existing_jobs_seen
                if stats.existing_jobs_seen
                else (len(records) - len(new_records))
            )

            summary_items = [
                ("Source URL", source_url),
                (
                    "Extraction Date (UTC)",
                    datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
                ),
                ("Total Jobs Extracted", len(records)),
                ("New Jobs Added Today", new_count),
                ("Existing Jobs Retained", existing_count),
                ("Small Companies (1-50 employees)", len(small_records)),
                ("Medium Companies (50-500 employees)", len(medium_records)),
                ("Large Companies (500+ employees)", len(large_records)),
                ("Unique Companies Found", stats.companies_discovered),
                ("Employee Counts Missing (N/A)", stats.missing_employee_counts),
                ("Pages Processed", stats.pages_processed),
            ]

            for row_idx, (key, value) in enumerate(summary_items, start=3):
                ws_summary.append([key, value])
                ws_summary.row_dimensions[row_idx].height = 22

                k_cell = ws_summary.cell(row=row_idx, column=1)
                k_cell.font = cls.SUMMARY_KEY_FONT
                k_cell.border = cls.THIN_BORDER
                k_cell.fill = PatternFill(
                    start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
                )
                k_cell.alignment = Alignment(horizontal="left", vertical="center")

                v_cell = ws_summary.cell(row=row_idx, column=2)
                v_cell.font = cls.DATA_FONT
                v_cell.border = cls.THIN_BORDER
                if key == "New Jobs Added Today" and value > 0:
                    v_cell.font = cls.NEW_STATUS_FONT
                    v_cell.fill = cls.NEW_STATUS_FILL
                if isinstance(value, int):
                    v_cell.alignment = Alignment(horizontal="right", vertical="center")
                    v_cell.number_format = "#,##0"
                else:
                    v_cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_summary.column_dimensions["A"].width = 38
            ws_summary.column_dimensions["B"].width = 50

            # Ensure parent directory exists
            os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

            wb.save(output_path)
            logger.info("Successfully exported %d records to %s", len(records), output_path)
            return os.path.abspath(output_path)

        except Exception as e:
            logger.error("Failed to generate Excel file: %s", e)
            raise ExcelExportError(f"Failed to generate Excel file: {e!s}") from e
